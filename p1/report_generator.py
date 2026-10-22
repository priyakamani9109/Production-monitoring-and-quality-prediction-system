"""
report_generator.py

Generates a structured Microsoft Word production report using python-docx,
based on the inputs and prediction results produced by the dashboard.
"""

import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def generate_report(inputs: dict, results: dict, insights: dict, recommendations: list) -> bytes:
    """Builds the Word report in-memory and returns the raw .docx bytes."""

    doc = Document()

    title = doc.add_heading("Production Monitoring Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("Supreme Nonwoven Industries Pvt. Ltd.")
    run.bold = True
    run.font.size = Pt(12)

    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_para.add_run(f"Generated on: {datetime.now().strftime('%d %B %Y, %I:%M %p')}")

    doc.add_paragraph()

    # ---- Production Details ----
    doc.add_heading("Production Details", level=1)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in inputs.items():
        row = table.add_row().cells
        row[0].text = str(label)
        row[1].text = str(value)

    doc.add_paragraph()

    # ---- Prediction Results ----
    doc.add_heading("Prediction Results", level=1)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in results.items():
        row = table.add_row().cells
        row[0].text = str(label)
        row[1].text = str(value)

    doc.add_paragraph()

    # ---- Production Insights ----
    doc.add_heading("Production Insights", level=1)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in insights.items():
        row = table.add_row().cells
        row[0].text = str(label)
        row[1].text = str(value)

    doc.add_paragraph()

    # ---- Recommendations ----
    doc.add_heading("Production Recommendations", level=1)
    if recommendations:
        for rec in recommendations:
            doc.add_paragraph(rec, style="List Bullet")
    else:
        doc.add_paragraph("No specific recommendations. Production conditions are within normal range.")

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Generated automatically by the Production Monitoring and Quality Prediction System")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
